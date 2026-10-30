"""Pencarian Rute Aman Menghindari Gangguan Jalan di Makassar (A*, UCS, BFS).

Jalankan GUI      : python main.py
Jalankan evaluasi : python main.py --eval

Data dibaca langsung dari berkas .xlsx di folder data/ (sheet Parameters, Nodes,
Edges, Hazard Locations, Test Scenarios). Tanpa .xlsx, program membaca CSV dengan
nama sheet yang sama (nodes.csv, edges.csv, ...).
Latar peta (opsional): data/map.png + kolom x_px, y_px pada sheet Nodes.
"""
import csv
import heapq
import math
import os
import random
import re
import sys
import time
from collections import deque

BASE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(BASE, "data")
MAP_IMAGE = os.path.join(DATA, "map.png")

GAMMA = 2.0            # dibaca dari sheet Parameters bila ada
PENALTY = {"tidak ada": 0, "ringan": 1, "berat": 2}
CLOSED = "ditutup"
KINDS = ["demonstrasi", "banjir", "kecelakaan"]
TIMING_RUNS = 100


# ----------------------------------------------------------------- data
def num(s):
    s = str(s).strip()
    if s.upper() == "INF":
        return math.inf
    return float(s.replace(",", "."))


_WB = {}


def find_workbook():
    if os.path.isdir(DATA):
        for f in sorted(os.listdir(DATA)):
            if f.lower().endswith(".xlsx") and not f.startswith("~"):
                return os.path.join(DATA, f)
    return None


def read_table(name):
    """Baca satu sheet (baris pertama = header) sebagai list of dict berisi string."""
    path = find_workbook()
    if path:
        try:
            from openpyxl import load_workbook
            if path not in _WB:
                _WB[path] = load_workbook(path, data_only=True)
            for ws in _WB[path].worksheets:
                if ws.title.strip().lower() == name.lower():
                    rows = list(ws.iter_rows(values_only=True))
                    head = [str(h).strip() if h is not None else "" for h in rows[0]]
                    return [dict(zip(head, ["" if v is None else str(v).strip() for v in r]))
                            for r in rows[1:] if sum(v is not None for v in r) >= 2]
            return []
        except ImportError:
            pass
    csv_path = os.path.join(DATA, name.lower().replace(" ", "_") + ".csv")
    if not os.path.exists(csv_path):
        return []
    with open(csv_path, newline="", encoding="utf-8-sig") as f:
        return [{k.strip(): (v or "").strip() for k, v in r.items()} for r in csv.DictReader(f)]


def get(row, prefix, default=""):
    """Ambil nilai kolom yang namanya diawali prefix (mis. 'x_km' untuk 'x_km (east)')."""
    for k, v in row.items():
        if k.lower().startswith(prefix):
            return v
    return default


class Graph:
    def __init__(self, nodes, edges):
        self.nodes, self.edges = nodes, edges
        self.adj = {n: [] for n in nodes}
        for e in edges.values():
            self.adj[e["a"]].append((e["b"], e["id"]))
            self.adj[e["b"]].append((e["a"], e["id"]))
        for n in self.adj:          # urutan tetangga tetap (sesuai id) agar hasil BFS bisa diulang
            self.adj[n].sort()


def load_graph():
    global GAMMA
    for r in read_table("Parameters"):
        if r.get("Parameter", "").lower().startswith("gamma"):
            GAMMA = num(r["Value"])
    nodes, edges = {}, {}
    for r in read_table("Nodes"):
        px, py = get(r, "x_px"), get(r, "y_px")
        nodes[r["node_id"]] = dict(
            id=r["node_id"], name=r["name"],
            kind="location" if r["type"].lower().startswith("start") else "junction",
            kx=num(get(r, "x_km")), ky=num(get(r, "y_km")),
            px=num(px) if px else None, py=num(py) if py else None,
            can_start=r.get("role", "").lower().startswith("start"))
    for r in read_table("Edges"):
        edges[r["edge_id"]] = dict(id=r["edge_id"], a=r["from"], b=r["to"],
                                   road=r["road_name"], dist=num(r["distance_km"]))
    if not nodes or not edges:
        sys.exit("Sheet Nodes / Edges tidak ditemukan. Taruh .xlsx dataset di folder data/.")
    for e in edges.values():
        if e["a"] not in nodes or e["b"] not in nodes:
            sys.exit(f"Ruas {e['id']} memakai simpul yang tidak ada di sheet Nodes.")
    return Graph(nodes, edges)


def straight_km(g, a, b):
    na, nb = g.nodes[a], g.nodes[b]
    return math.hypot(na["kx"] - nb["kx"], na["ky"] - nb["ky"])


def check_heuristic(g):
    """Heuristik admissible bila garis lurus <= jarak jalan pada setiap ruas."""
    return [e["id"] for e in g.edges.values()
            if straight_km(g, e["a"], e["b"]) > e["dist"] + 1e-9]


def hazard_defaults():
    """Sheet Hazard Locations -> {edge_id: (jenis, tingkat bawaan)}."""
    out = {}
    for r in read_table("Hazard Locations"):
        lv = r.get("default level", "").lower()
        level = CLOSED if lv.startswith("ditutup") else "berat" if lv.startswith("berat") else None
        if level:
            for e in re.findall(r"E\d+", r.get("graph edges", "")):
                out[e] = (r["hazard"].lower().split()[0], level)
    return out


def load_scenarios():
    """Sheet Test Scenarios -> list skenario dengan gangguan {edge: (jenis, tingkat)}."""
    out, defaults = [], None
    for r in read_table("Test Scenarios"):
        edges = re.findall(r"E\d+", r.get("hazard_edges", ""))
        level, htype = r.get("level", "-").lower(), r.get("hazard_type", "-").lower()
        dis = {}
        if edges and "+" in level + htype:      # gabungan: tingkat per ruas dari Hazard Locations
            defaults = defaults or hazard_defaults()
            dis = {e: defaults.get(e, ("demonstrasi", CLOSED)) for e in edges}
        elif edges:
            dis = {e: (htype.split()[0], CLOSED if level.startswith("ditutup") else level)
                   for e in edges}
        out.append(dict(id=r["ID"], name=r["scenario"], start=r["start"], goal=r["goal"],
                        dis=dis, expected=r.get("cost with hazard", "")))
    return out


# --------------------------------------------------------------- search
def edge_cost(g, eid, dis):
    """cost(e) = d(e) + gamma * p(e); ruas ditutup -> None (dihapus dari graf)."""
    level = dis.get(eid, (None, "tidak ada"))[1]
    if level == CLOSED:
        return None
    return g.edges[eid]["dist"] + GAMMA * PENALTY[level]


def finish(g, parent, goal, dis, expanded):
    nodes, eids, n = [goal], [], goal
    while n in parent:
        n, eid = parent[n]
        nodes.append(n)
        eids.append(eid)
    nodes.reverse()
    eids.reverse()
    return dict(found=True, nodes=nodes, edges=eids, expanded=expanded,
                dist=sum(g.edges[i]["dist"] for i in eids),
                cost=sum(edge_cost(g, i, dis) for i in eids))


def not_found(expanded):
    return dict(found=False, nodes=[], edges=[], expanded=expanded, dist=math.inf, cost=math.inf)


def best_first(g, start, goal, dis, use_h):
    """A* bila use_h=True (f = g + h), UCS/Dijkstra bila False (f = g)."""
    def h(n):
        return straight_km(g, n, goal) if use_h else 0.0

    tie = 0
    frontier = [(h(start), 0, start)]
    best, parent, done, expanded = {start: 0.0}, {}, set(), 0
    while frontier:
        _, _, n = heapq.heappop(frontier)
        if n in done:
            continue
        done.add(n)
        expanded += 1
        if n == goal:
            return finish(g, parent, goal, dis, expanded)
        for m, eid in g.adj[n]:
            c = edge_cost(g, eid, dis)
            if c is None:
                continue
            ng = best[n] + c
            if ng < best.get(m, math.inf):
                best[m] = ng
                parent[m] = (n, eid)
                tie += 1
                heapq.heappush(frontier, (ng + h(m), tie, m))
    return not_found(expanded)


def bfs(g, start, goal, dis):
    """Rute dengan jumlah ruas paling sedikit, bobot diabaikan (ruas ditutup tetap dihindari)."""
    queue, seen, parent, expanded = deque([start]), {start}, {}, 0
    while queue:
        n = queue.popleft()
        expanded += 1
        if n == goal:
            return finish(g, parent, goal, dis, expanded)
        for m, eid in g.adj[n]:
            if m in seen or edge_cost(g, eid, dis) is None:
                continue
            seen.add(m)
            parent[m] = (n, eid)
            queue.append(m)
    return not_found(expanded)


ALGOS = {
    "A*": lambda g, s, t, d: best_first(g, s, t, d, True),
    "UCS": lambda g, s, t, d: best_first(g, s, t, d, False),
    "BFS": bfs,
}


def run_algo(name, g, s, t, dis, runs=TIMING_RUNS):
    fn = ALGOS[name]
    t0 = time.perf_counter()
    for _ in range(runs):
        res = fn(g, s, t, dis)
    res["ms"] = (time.perf_counter() - t0) * 1000 / runs
    return res


def fc(x):
    return "INF" if x == math.inf else f"{x:.2f}"


# ----------------------------------------------------------- evaluation
def evaluate(g):
    scenarios = load_scenarios()
    if not scenarios:
        print("Sheet Test Scenarios kosong atau tidak ada.")
        return
    print(f"{'ID':<4}{'Skenario':<30}{'Harap':>7}{'UCS':>8}{'A*':>8}{'BFS':>8}"
          f"{'Opt':>5}{'Ubah':>6}{'Aman':>6}{'Eks B/U/A':>12}{'ms A*':>8}")
    n_opt = n_safe = n_exp = n_exp_ok = bfs_opt = 0
    exp_sum = dict.fromkeys(ALGOS, 0)
    ms_sum = dict.fromkeys(ALGOS, 0.0)
    for sc in scenarios:
        s, t, dis = sc["start"], sc["goal"], sc["dis"]
        res = {a: run_algo(a, g, s, t, dis) for a in ALGOS}
        base = ALGOS["A*"](g, s, t, {})
        same = lambda x, y: x == y or abs(x - y) < 1e-9
        opt = same(res["A*"]["cost"], res["UCS"]["cost"])
        changed = res["A*"]["nodes"] != base["nodes"]
        safe = all(dis.get(e, (0, ""))[1] != CLOSED for e in res["A*"]["edges"])
        mark = "-"
        if sc["expected"]:
            ok = same(num(sc["expected"]), res["UCS"]["cost"]) or \
                abs(num(sc["expected"]) - res["UCS"]["cost"]) < 0.011
            mark = "OK" if ok else "X"
            n_exp += 1
            n_exp_ok += ok
        n_opt += opt
        n_safe += safe
        bfs_opt += same(res["BFS"]["cost"], res["UCS"]["cost"])
        for a in ALGOS:
            exp_sum[a] += res[a]["expanded"]
            ms_sum[a] += res[a]["ms"]
        print(f"{sc['id']:<4}{sc['name'][:28]:<30}{mark:>7}{fc(res['UCS']['cost']):>8}"
              f"{fc(res['A*']['cost']):>8}{fc(res['BFS']['cost']):>8}"
              f"{'Y' if opt else 'N':>5}{'Y' if changed else 'N':>6}{'Y' if safe else 'N':>6}"
              f"{res['BFS']['expanded']:>5}/{res['UCS']['expanded']}/{res['A*']['expanded']:<3}"
              f"{res['A*']['ms']:>8.3f}")
    n = len(scenarios)
    print(f"\nOptimalitas A* terhadap UCS : {n_opt}/{n} ({100 * n_opt / n:.0f}%)")
    print(f"BFS optimal                 : {bfs_opt}/{n}")
    print(f"Rute A* tidak melewati ruas ditutup: {n_safe}/{n}")
    if n_exp:
        print(f"Cocok dengan biaya di sheet Test Scenarios: {n_exp_ok}/{n_exp}")
    for a in ALGOS:
        print(f"{a:<4} rata-rata simpul diekspansi {exp_sum[a] / n:5.1f} | rata-rata waktu {ms_sum[a] / n:.3f} ms")
    simulate_accidents(g)


def simulate_accidents(g, n=10, seed=42):
    """Kecelakaan acak (ruas mana pun, ringan/berat) pada rute UC -> MP, seed tetap."""
    s, t = "UC", "MP"
    if s not in g.nodes or t not in g.nodes:
        return
    rnd = random.Random(seed)
    edge_ids = list(g.edges)
    base = ALGOS["A*"](g, s, t, {})
    print(f"\nSimulasi kecelakaan acak UC -> MP (seed {seed}), biaya dasar {fc(base['cost'])}:")
    ok = 0
    for k in range(1, n + 1):
        eid, level = rnd.choice(edge_ids), rnd.choice(["ringan", "berat"])
        dis = {eid: ("kecelakaan", level)}
        a, u = ALGOS["A*"](g, s, t, dis), ALGOS["UCS"](g, s, t, dis)
        same = a["cost"] == u["cost"] or abs(a["cost"] - u["cost"]) < 1e-9
        ok += same
        print(f"  {k:>2}. {eid} {level:<7} A*={fc(a['cost'])} UCS={fc(u['cost'])} "
              f"{'berubah' if a['nodes'] != base['nodes'] else 'tetap':<8}{'OK' if same else 'GAGAL'}")
    print(f"  A* = UCS pada {ok}/{n} simulasi")


# ------------------------------------------------------------------ GUI
LEVEL_COLOR = {"ringan": "#f4c20d", "berat": "#f57c00", "ditutup": "#d32f2f"}


def seg_dist(px, py, ax, ay, bx, by):
    dx, dy = bx - ax, by - ay
    L = dx * dx + dy * dy
    t = 0 if L == 0 else max(0, min(1, ((px - ax) * dx + (py - ay) * dy) / L))
    return math.hypot(px - (ax + t * dx), py - (ay + t * dy))


def layout(g, W, H, use_px, margin=60):
    """Posisi gambar tiap simpul: x_px/y_px bila ada, selain itu skema dari x_km/y_km."""
    if use_px:
        for n in g.nodes.values():
            n["sx"], n["sy"] = n["px"], n["py"]
        return
    xs, ys = [n["kx"] for n in g.nodes.values()], [n["ky"] for n in g.nodes.values()]
    k = min((W - 2 * margin) / (max(xs) - min(xs) or 1), (H - 2 * margin) / (max(ys) - min(ys) or 1))
    for n in g.nodes.values():
        n["sx"] = margin + (n["kx"] - min(xs)) * k
        n["sy"] = H - margin - (n["ky"] - min(ys)) * k


def run_gui(g, heuristic_note):
    import tkinter as tk
    from tkinter import ttk

    root = tk.Tk()
    root.title("Rute Aman Makassar - A*")
    dis, state = {}, {"results": {}}

    has_px = all(n["px"] is not None and n["py"] is not None for n in g.nodes.values())
    img = tk.PhotoImage(file=MAP_IMAGE) if has_px and os.path.exists(MAP_IMAGE) else None
    W, H = (img.width(), img.height()) if img else (900, 650)
    layout(g, W, H, has_px)
    canvas = tk.Canvas(root, width=W, height=H, bg="white")
    canvas.grid(row=0, column=0)
    panel = ttk.Frame(root, padding=8)
    panel.grid(row=0, column=1, sticky="ns")

    name_to_id = {n["name"]: i for i, n in g.nodes.items()}
    starts = [n["name"] for n in g.nodes.values() if n["can_start"]]
    goals = [n["name"] for n in g.nodes.values() if n["kind"] == "location"]
    start_var, goal_var, algo_var = tk.StringVar(), tk.StringVar(), tk.StringVar(value="A*")

    ttk.Label(panel, text="Titik awal").pack(anchor="w")
    cb_s = ttk.Combobox(panel, textvariable=start_var, values=starts, state="readonly", width=34)
    cb_s.pack(anchor="w")
    ttk.Label(panel, text="Tujuan").pack(anchor="w", pady=(6, 0))
    cb_g = ttk.Combobox(panel, textvariable=goal_var, values=goals, state="readonly", width=34)
    cb_g.pack(anchor="w")
    ttk.Label(panel, text="Rute yang digambar").pack(anchor="w", pady=(6, 0))
    row = ttk.Frame(panel)
    row.pack(anchor="w")
    for a in ALGOS:
        ttk.Radiobutton(row, text=a, value=a, variable=algo_var,
                        command=lambda: refresh()).pack(side="left", padx=(0, 8))
    ttk.Button(panel, text="Hapus semua gangguan",
               command=lambda: (dis.clear(), search())).pack(anchor="w", pady=6)
    ttk.Label(panel, text="Klik garis jalan pada peta untuk menandai gangguan.",
              wraplength=320).pack(anchor="w")
    text = tk.Text(panel, width=46, height=30, font=("Courier", 9), state="disabled")
    text.pack(pady=6)

    def show(msg):
        text.config(state="normal")
        text.delete("1.0", "end")
        text.insert("end", msg)
        text.config(state="disabled")

    def draw():
        canvas.delete("all")
        if img:
            canvas.create_image(0, 0, image=img, anchor="nw")
        shown = state["results"].get(algo_var.get())
        on_route = set(shown["edges"]) if shown and shown["found"] else set()
        for e in g.edges.values():
            a, b = g.nodes[e["a"]], g.nodes[e["b"]]
            canvas.create_line(a["sx"], a["sy"], b["sx"], b["sy"], fill="#666", width=2)
        for eid in on_route:
            e = g.edges[eid]
            a, b = g.nodes[e["a"]], g.nodes[e["b"]]
            canvas.create_line(a["sx"], a["sy"], b["sx"], b["sy"], fill="#1565c0", width=6)
        for eid, (kind, level) in dis.items():
            e = g.edges[eid]
            a, b = g.nodes[e["a"]], g.nodes[e["b"]]
            if eid not in on_route:
                canvas.create_line(a["sx"], a["sy"], b["sx"], b["sy"], fill=LEVEL_COLOR[level],
                                   width=4, dash=(6, 4) if level == CLOSED else None)
            mx, my = (a["sx"] + b["sx"]) / 2, (a["sy"] + b["sy"]) / 2
            canvas.create_text(mx, my - 10, text=f"{kind[:5]}-{level}", fill=LEVEL_COLOR[level],
                               font=("Helvetica", 8, "bold"))
        s, t = name_to_id.get(start_var.get()), name_to_id.get(goal_var.get())
        for i, n in g.nodes.items():
            r = 9 if n["kind"] == "location" else 5
            fill = "#2e7d32" if i == s else "#6a1b9a" if i == t else "white"
            canvas.create_oval(n["sx"] - r, n["sy"] - r, n["sx"] + r, n["sy"] + r,
                               fill=fill, outline="black", width=2)
            if n["kind"] == "location":
                canvas.create_text(n["sx"], n["sy"] - 16, text=n["id"],
                                   font=("Helvetica", 10, "bold"))

    def show_results():
        res = state["results"]
        if not res:
            return
        shown = res[algo_var.get()]
        out = []
        if not shown["found"]:
            out.append("Tujuan tidak dapat dijangkau:\nsemua rute terhalang ruas yang ditutup.")
        else:
            out.append(f"Rute ({algo_var.get()}):")
            for i, eid in enumerate(shown["edges"]):
                e = g.edges[eid]
                a, b = g.nodes[shown["nodes"][i]]["name"], g.nodes[shown["nodes"][i + 1]]["name"]
                d = dis.get(eid)
                note = f" [{d[0]}, {d[1]}]" if d else ""
                out.append(f"{i + 1}. {a} -> {b}\n   {e['road']} ({e['dist']:.2f} km){note}")
            out.append(f"\nTotal jarak : {shown['dist']:.2f} km\nTotal biaya : {shown['cost']:.2f}")
        out.append("\nPerbandingan algoritma")
        out.append(f"{'Algo':<5}{'Biaya':>8}{'Jarak':>8}{'Ekspansi':>10}{'ms':>8}")
        for a, r in res.items():
            out.append(f"{a:<5}{fc(r['cost']):>8}{fc(r['dist']):>8}{r['expanded']:>10}{r['ms']:>8.3f}")
        show("\n".join(out))

    def refresh():
        draw()
        show_results()

    def search(*_):
        s, t = name_to_id.get(start_var.get()), name_to_id.get(goal_var.get())
        state["results"] = {}
        if not s or not t:
            draw()
            show("Pilih titik awal dan tujuan.\n\n" + heuristic_note)
            return
        if s == t:
            draw()
            show("Titik awal dan tujuan tidak boleh sama.")
            return
        state["results"] = {a: run_algo(a, g, s, t, dis) for a in ALGOS}
        refresh()

    def edge_dialog(eid):
        e = g.edges[eid]
        win = tk.Toplevel(root)
        win.title("Gangguan ruas")
        win.transient(root)
        win.grab_set()
        ttk.Label(win, padding=8, text=f"{e['road']}\n{g.nodes[e['a']]['name']} - "
                                       f"{g.nodes[e['b']]['name']} ({e['dist']:.2f} km)").grid(
            row=0, column=0, columnspan=2)
        cur = dis.get(eid, (KINDS[0], "ringan"))
        kind, level = tk.StringVar(value=cur[0]), tk.StringVar(value=cur[1])
        ttk.Label(win, text="Jenis").grid(row=1, column=0, padx=8, sticky="w")
        ttk.Combobox(win, textvariable=kind, values=KINDS, state="readonly").grid(row=1, column=1, padx=8)
        ttk.Label(win, text="Tingkat").grid(row=2, column=0, padx=8, sticky="w")
        ttk.Combobox(win, textvariable=level, values=["ringan", "berat", CLOSED],
                     state="readonly").grid(row=2, column=1, padx=8)

        def apply():
            dis[eid] = (kind.get(), level.get())
            win.destroy()
            search()

        def clear():
            dis.pop(eid, None)
            win.destroy()
            search()

        bar = ttk.Frame(win, padding=8)
        bar.grid(row=3, column=0, columnspan=2)
        ttk.Button(bar, text="Terapkan", command=apply).pack(side="left", padx=4)
        ttk.Button(bar, text="Hapus gangguan", command=clear).pack(side="left", padx=4)

    def on_click(ev):
        best = None
        for e in g.edges.values():
            a, b = g.nodes[e["a"]], g.nodes[e["b"]]
            d = seg_dist(ev.x, ev.y, a["sx"], a["sy"], b["sx"], b["sy"])
            if d < 8 and (best is None or d < best[0]):
                best = (d, e["id"])
        if best:
            edge_dialog(best[1])

    canvas.bind("<Button-1>", on_click)
    cb_s.bind("<<ComboboxSelected>>", search)
    cb_g.bind("<<ComboboxSelected>>", search)
    draw()
    show("Pilih titik awal dan tujuan.\n\n" + heuristic_note)
    root.mainloop()


# ----------------------------------------------------------------- main
if __name__ == "__main__":
    graph = load_graph()
    bad = check_heuristic(graph)
    if bad:
        note = f"PERINGATAN: heuristik tidak admissible pada ruas {', '.join(bad)}."
    else:
        note = f"Heuristik admissible: lolos pada {len(graph.edges)} ruas."
    print(f"{len(graph.nodes)} simpul, {len(graph.edges)} ruas, gamma = {GAMMA:g}. {note}")
    if "--eval" in sys.argv:
        evaluate(graph)
    else:
        run_gui(graph, note)